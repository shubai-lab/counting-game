# -*- coding: utf-8 -*-
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

with open('excel_content_utf8.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

generals = data['content']['武将名单']['data'][2:]
name_to_quality = {row[1]: row[2] for row in generals if row[0] and row[1] and row[2]}

# 限定武将池武将及其来源分区
limited_generals = [
    ('关索', '限定SSS'),
    ('赵襄', '限定SSS'),
    ('鲍三娘', '限定SSS'),
    ('徐荣', '限定SSS'),
    ('曹婴', '限定SSS'),
    ('张琪瑛', '限定SSS'),
    ('曹纯', '限定SSS'),
    ('花鬘', '限定SSS'),
    ('刘宏', '限定(高级)'),
    ('刘永', '限定(高级)'),
    ('张勋', '限定(高级)'),
    ('麴义', '限定(高级)'),
    ('祢衡', '限定(高级)'),
    ('胡金定', '限定(高级)'),
    ('皇甫嵩', '限定(普通)'),
    ('赵昂', '限定(普通)'),
    ('何进', '限定(普通)'),
    ('严夫人', '限定(普通)'),
    ('陆凯', '限定(普通)'),
    ('管辂', '限定(普通)'),
    ('王昶', '限定(普通)'),
]

# 创建数据
rows = []
for name, old_pool in limited_generals:
    new_quality = name_to_quality.get(name, '未找到')
    rows.append({
        '武将名称': name,
        '原分区': old_pool,
        '新品质': new_quality
    })

df = pd.DataFrame(rows)

# 创建Excel
wb = Workbook()
ws = wb.active
ws.title = '限定武将池品质对照'

# 样式
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=12)
header_alignment = Alignment(horizontal='center', vertical='center')

quality_fills = {
    '限定': PatternFill(start_color='FFD700', end_color='FFD700', fill_type='solid'),  # 金色
    '传说': PatternFill(start_color='9400D3', end_color='9400D3', fill_type='solid'),  # 紫色
    '史诗': PatternFill(start_color='4169E1', end_color='4169E1', fill_type='solid'),  # 蓝色
    '稀有': PatternFill(start_color='32CD32', end_color='32CD32', fill_type='solid'),  # 绿色
    '普通': PatternFill(start_color='A9A9A9', end_color='A9A9A9', fill_type='solid'),  # 灰色
}
quality_fonts = {
    '限定': Font(bold=True, color='000000', size=11),
    '传说': Font(bold=True, color='FFFFFF', size=11),
    '史诗': Font(bold=True, color='FFFFFF', size=11),
    '稀有': Font(bold=True, color='000000', size=11),
    '普通': Font(bold=True, color='000000', size=11),
}

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 写入表头
headers = ['武将名称', '原分区', '新品质']
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border

# 写入数据
for row_idx, row_data in enumerate(rows, 2):
    ws.cell(row=row_idx, column=1, value=row_data['武将名称']).border = thin_border
    ws.cell(row=row_idx, column=2, value=row_data['原分区']).border = thin_border

    cell = ws.cell(row=row_idx, column=3, value=row_data['新品质'])
    cell.border = thin_border
    cell.alignment = header_alignment

    quality = row_data['新品质']
    if quality in quality_fills:
        cell.fill = quality_fills[quality]
        cell.font = quality_fonts[quality]

# 设置列宽
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 15

# 添加统计Sheet
ws2 = wb.create_sheet('统计')

# 统计
quality_count = {}
for name, _ in limited_generals:
    q = name_to_quality.get(name, '未知')
    quality_count[q] = quality_count.get(q, 0) + 1

# 写入统计
stats_headers = ['新品质', '人数']
for col, header in enumerate(stats_headers, 1):
    cell = ws2.cell(row=1, column=col, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_alignment
    cell.border = thin_border

row_idx = 2
for q in ['限定', '传说', '史诗', '稀有', '普通']:
    if q in quality_count:
        ws2.cell(row=row_idx, column=1, value=q).border = thin_border
        cell = ws2.cell(row=row_idx, column=2, value=quality_count[q])
        cell.border = thin_border
        cell.alignment = header_alignment
        if q in quality_fills:
            cell.fill = quality_fills[q]
            cell.font = quality_fonts[q]
        row_idx += 1

# 合计
ws2.cell(row=row_idx, column=1, value='合计').border = thin_border
ws2.cell(row=row_idx, column=1).font = Font(bold=True, size=12)
ws2.cell(row=row_idx, column=2, value=sum(quality_count.values())).border = thin_border
ws2.cell(row=row_idx, column=2).font = Font(bold=True, size=12)

ws2.column_dimensions['A'].width = 15
ws2.column_dimensions['B'].width = 15

# 保存
output_path = '限定武将池_品质对照表.xlsx'
wb.save(output_path)
print(f'Excel文件已生成: {output_path}')
