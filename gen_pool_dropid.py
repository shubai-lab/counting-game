# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import json

# 读取Excel
wb = load_workbook('将池武将_品质对照表.xlsx')
ws = wb['Sheet1']

# DropGroupID与将池名称对照
dropid_to_pool = {
    601: '祈福SSS武将池',
    602: '限定SSS武将池',
    603: '高级SS武将池',
    604: 'SS级武将池',
    605: 'S级武将池',
}

# DropGroupID与新品质对照
dropid_to_quality = {
    601: '传说',  # 祈福SSS -> 传说
    602: '限定',  # 限定SSS -> 限定
    603: '史诗',  # 高级SS -> 史诗
    604: '史诗',  # SS级 -> 史诗
    605: '稀有',  # S级 -> 稀有
}

# 读取所有武将数据
all_generals = []
for r in range(5, ws.max_row + 1):
    conf_id = ws.cell(r, 1).value
    drop_id = ws.cell(r, 2).value
    name = ws.cell(r, 3).value
    if conf_id and drop_id and name and isinstance(drop_id, int):
        pool_name = dropid_to_pool.get(drop_id, f'其他({drop_id})')
        quality = dropid_to_quality.get(drop_id, '未知')
        all_generals.append({
            'conf_id': conf_id,
            'drop_id': drop_id,
            'name': name,
            'pool': pool_name,
            'quality': quality
        })

print(f'总武将数: {len(all_generals)}')

# 按DropID统计
from collections import Counter
drop_stats = Counter(g['drop_id'] for g in all_generals)
print('\n按DropGroupID统计:')
for drop_id in sorted(drop_stats.keys()):
    print(f'  {drop_id} ({dropid_to_pool.get(drop_id, "其他")}): {drop_stats[drop_id]}人')

# 按品质统计
quality_stats = Counter(g['quality'] for g in all_generals)
print('\n按新品质统计:')
for q in ['限定', '传说', '史诗', '稀有', '普通']:
    if q in quality_stats:
        print(f'  {q}: {quality_stats[q]}人')

# 输出前20个示例
print('\n前20个武将示例:')
for g in all_generals[:20]:
    print(f'  {g["name"]} | {g["pool"]} | {g["quality"]}')
