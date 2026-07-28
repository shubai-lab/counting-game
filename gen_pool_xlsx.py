# -*- coding: utf-8 -*-
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

with open('excel_content_utf8.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

generals = data['content']['武将名单']['data'][2:]
name_to_quality = {row[1]: row[2] for row in generals if row[0] and row[1] and row[2]}

# 网页中的所有武将
web_pools = {
    '祈福SSS': ['司马徽', '兀突骨', '庞德公', '王烈', '沙摩柯', '陆郁生'],
    '限定SSS': ['关索', '赵襄', '鲍三娘', '徐荣', '曹婴', '张琪瑛', '曹纯', '花鬘'],
    '限定(高级)': ['刘宏', '刘永', '张勋', '麴义', '祢衡', '胡金定'],
    '限定(普通)': ['皇甫嵩', '赵昂', '何进', '严夫人', '陆凯', '管辂', '王昶'],
    '活动武将池': ['张让', '王桃', '陈珪', '辛宪英', '黄祖', '韩馥', '王悦', '胡班', '蒲元', '杨仪', '孙鲁育', '邓芝', '田畴', '王荣', '南华老仙', '华歆', '冯熙', '轲比能', '雷铜', '吕玲绮', '陶谦', '朱灵', '雷薄', '王双', '张嫙', '张邈', '吉平', '尹夫人', '刘琦', '潘淑', '曹安民', '刘巴', '王威', '董承', '滕公主', '何晏', '冯方', '穆顺', '张虎', '管亥', '唐姬', '滕胤', '吕旷吕翔', '留赞', '诸葛尚', '杜夫人', '羊祜', '刘辩', '杨婉', '荀谌', '卞喜'],
    'SS武将池': ['黄忠', '邓艾', '郝昭', '虞翻', '步骘', '戏志才', '蒋琬费祎', '魏延', '姜维', '曹植', '伏皇后', '曹昂', 'SP孙尚香', '李傕郭汜', '曹仁', '刘禅', '陈宫', '李儒', '刘协', 'SP蔡文姬', '董白', '小乔', '孙策', '关兴张苞', '张松', '曹节', '李严', '赵忠', '周泰', '张昭张纮', '步练师', '沮授', '董允', '糜竺', '崔琰毛玠', '张角', '蔡文姬', '荀攸', '曹叡', '诸葛瑾', '关银屏', '朱儁', '荀彧', '陆绩', '王异', '钟繇', '李典', '徐氏', '曹丕', '王基', '钟会', '夏侯氏', '刘焉', '吕虔', '贾诩', '许攸', '曹冲', '朱治', '张鲁', '马腾', '鲁肃', '诸葛瞻', '满宠', '郭图逢纪', '秦宓', '孔融', '张郃', '周妃', '刘封', '吴苋', '薛综', '田丰'],
    'S武将池': ['于吉', '王平', '凌统', '陈群', '诸葛诞', 'SP黄月英', '臧霸', '庞统', '孙亮', '徐盛', '吴懿', '严畯', '周鲂', '糜夫人', '卧龙诸葛', '蒯良蒯越', '张春华', '周仓', '司马朗', '贺齐', '陈武董袭', '太史慈', '陈到', '高顺', '孙鲁班', '杜畿', '文聘', '张任', '庞德', '陆抗', '马岱', '朱桓', '孙乾', '吕岱', '曹洪', '颜良文丑', '毌丘俭', '程普', '顾雍', '潘濬', '刘繇', '蒋钦', '袁绍', '袁术', '韩当', '蔡夫人', '郭皇后', '马云騄', '卞夫人', '徐晃', '张绣', '刘表', '曹休', '王粲', '祖茂', '孙坚', '法正', '关平', '刘谌', '蔡邕', '蹋顿', '董卓', '马谡', '简雍', '孙休', 'SP徐庶', '乐进', '祝融', '徐庶', '潘璋马忠', '公孙渊', 'SP庞统', '纪灵', '孟获', '吴国太', '韩浩史涣', '孙登', 'SP姜维', '何太后'],
}

wb = Workbook()
ws = wb.active
ws.title = '将池品质对照'

# 样式
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=11)
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

# 表头
for col, title in enumerate(['武将名称', '原分区', '新品质'], 1):
    cell = ws.cell(row=1, column=col, value=title)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = thin_border
    cell.alignment = Alignment(horizontal='center')

# 数据
row_idx = 2
all_data = []
for pool_name, names in web_pools.items():
    for name in names:
        new_q = name_to_quality.get(name, '未知')
        all_data.append((name, pool_name, new_q))
        ws.cell(row=row_idx, column=1, value=name).border = thin_border
        ws.cell(row=row_idx, column=2, value=pool_name).border = thin_border
        ws.cell(row=row_idx, column=3, value=new_q).border = thin_border
        ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal='center')
        row_idx += 1

# 列宽
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 10

# 统计Sheet
ws2 = wb.create_sheet('统计')
ws2.cell(row=1, column=1, value='原分区').fill = header_fill
ws2.cell(row=1, column=1, value='原分区').font = header_font
ws2.cell(row=1, column=2, value='人数').fill = header_fill
ws2.cell(row=1, column=2, value='人数').font = header_font

r = 2
for pool, names in web_pools.items():
    ws2.cell(row=r, column=1, value=pool).border = thin_border
    ws2.cell(row=r, column=2, value=len(names)).border = thin_border
    ws2.cell(row=r, column=2).alignment = Alignment(horizontal='center')
    r += 1

total = sum(len(v) for v in web_pools.values())
ws2.cell(row=r, column=1, value='合计').border = thin_border
ws2.cell(row=r, column=1).font = Font(bold=True)
ws2.cell(row=r, column=2, value=total).border = thin_border
ws2.cell(row=r, column=2).font = Font(bold=True)
ws2.cell(row=r, column=2).alignment = Alignment(horizontal='center')

ws2.column_dimensions['A'].width = 15
ws2.column_dimensions['B'].width = 10

wb.save('pool_quality_final.xlsx')

# 输出验证
import sys
sys.stdout.reconfigure(encoding='utf-8')
print(f'总武将数: {len(all_data)}')
print(f'\n新品质统计:')
from collections import Counter
qc = Counter(d[2] for d in all_data)
for q in ['限定', '传说', '史诗', '稀有', '普通']:
    if q in qc:
        print(f'  {q}: {qc[q]}')
