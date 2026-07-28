# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import json

# 读取武将品质数据
with open('excel_content_utf8.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

generals = data['content']['武将名单']['data'][2:]
name_to_quality = {row[1]: row[2] for row in generals if row[0] and row[1] and row[2]}

# 读取Excel
wb = load_workbook('将池武将_品质对照表.xlsx')
ws = wb['Sheet1']

# 品质颜色
quality_fills = {
    '限定': PatternFill(start_color='FFD700', end_color='FFD700', fill_type='solid'),
    '传说': PatternFill(start_color='9400D3', end_color='9400D3', fill_type='solid'),
    '史诗': PatternFill(start_color='4169E1', end_color='4169E1', fill_type='solid'),
    '稀有': PatternFill(start_color='32CD32', end_color='32CD32', fill_type='solid'),
    '普通': PatternFill(start_color='A9A9A9', end_color='A9A9A9', fill_type='solid'),
}
quality_fonts = {
    '限定': Font(bold=True, color='000000', size=10),
    '传说': Font(bold=True, color='FFFFFF', size=10),
    '史诗': Font(bold=True, color='FFFFFF', size=10),
    '稀有': Font(bold=True, color='000000', size=10),
    '普通': Font(bold=True, color='000000', size=10),
}
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

# 统计
from collections import Counter
quality_count = Counter()
item_count = 0

# 添加新列标题
ws.cell(row=4, column=4, value='新品质')
ws.cell(row=1, column=4, value='int32')
ws.cell(row=2, column=4, value='string')

# 填充品质数据
for r in range(5, ws.max_row + 1):
    name = ws.cell(r, 3).value
    if name:
        quality = name_to_quality.get(name, '')
        cell = ws.cell(row=r, column=4, value=quality)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal='center')
        if quality in quality_fills:
            cell.fill = quality_fills[quality]
            cell.font = quality_fonts[quality]
            quality_count[quality] += 1
            item_count += 1
        elif quality == '':
            item_count += 1  # 非武将道具也计入

# 设置列宽
ws.column_dimensions['D'].width = 12

# 保存
output_path = '将池武将_品质对照表_新版.xlsx'
wb.save(output_path)

print(f'Excel已生成: {output_path}')
print(f'\n品质统计:')
for q in ['限定', '传说', '史诗', '稀有', '普通']:
    if q in quality_count:
        print(f'  {q}: {quality_count[q]}')
print(f'  总计: {item_count}')
