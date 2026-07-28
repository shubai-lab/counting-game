# -*- coding: utf-8 -*-
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

with open('excel_content_utf8.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

generals = data['content']['武将名单']['data'][2:]
name_to_quality = {row[1]: row[2] for row in generals if row[0] and row[1] and row[2]}

web_pools = {
    '祈福SSS': ['司马徽', '兀突骨', '庞德公', '王烈', '沙摩柯', '陆郁生'],
    '限定SSS': ['关索', '赵襄', '鲍三娘', '徐荣', '曹婴', '张琪瑛', '曹纯', '花鬘'],
    '限定(高级)': ['刘宏', '刘永', '张勋', '麴义', '祢衡', '胡金定'],
    '限定(普通)': ['皇甫嵩', '赵昂', '何进', '严夫人', '陆凯', '管辂', '王昶'],
}

wb = Workbook()
ws = wb.active
ws.title = '将池品质对照'

# 样式
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF', size=11)

thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 表头
ws.cell(row=1, column=1, value='武将名称').fill = header_fill
ws.cell(row=1, column=1, value='武将名称').font = header_font
ws.cell(row=1, column=1, value='武将名称').border = thin_border

ws.cell(row=1, column=2, value='原分区').fill = header_fill
ws.cell(row=1, column=2, value='原分区').font = header_font
ws.cell(row=1, column=2, value='原分区').border = thin_border

ws.cell(row=1, column=3, value='新品质').fill = header_fill
ws.cell(row=1, column=3, value='新品质').font = header_font
ws.cell(row=1, column=3, value='新品质').border = thin_border

# 数据
row_idx = 2
for pool_name, names in web_pools.items():
    for name in names:
        new_q = name_to_quality.get(name, '未知')
        ws.cell(row=row_idx, column=1, value=name).border = thin_border
        ws.cell(row=row_idx, column=2, value=pool_name).border = thin_border
        ws.cell(row=row_idx, column=3, value=new_q).border = thin_border
        row_idx += 1

ws.column_dimensions['A'].width = 12
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 10

wb.save('pool_quality.xlsx')
print('Done')
