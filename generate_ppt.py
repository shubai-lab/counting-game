# -*- coding: utf-8 -*-
"""
游戏策划需求 Skill PPT 生成脚本
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml

# 配色方案
COLORS = {
    'primary': RGBColor(0x1E, 0x3A, 0x5F),      # 深蓝 #1E3A5F
    'secondary': RGBColor(0x34, 0x98, 0xDB),    # 亮蓝 #3498DB
    'accent': RGBColor(0xE7, 0x4C, 0x3C),       # 红 #E74C3C
    'background': RGBColor(0xF5, 0xF7, 0xFA),   # 浅灰 #F5F7FA
    'text': RGBColor(0x2C, 0x3E, 0x50),         # 深灰 #2C3E50
    'white': RGBColor(0xFF, 0xFF, 0xFF),        # 白色
    'light_blue': RGBColor(0xE8, 0xF4, 0xFC),    # 浅蓝背景
}

def set_shape_fill(shape, color):
    """设置形状填充颜色"""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color

def set_shape_no_fill(shape):
    """设置形状无填充"""
    shape.fill.background()

def set_text_frame(shape, text, font_size=14, font_color=None, bold=False, align=PP_ALIGN.LEFT):
    """设置文本框内容"""
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    if font_color:
        p.font.color.rgb = font_color
    p.alignment = align

def add_text_box(slide, left, top, width, height, text, font_size=14, font_color=None, bold=False, align=PP_ALIGN.LEFT):
    """添加文本框"""
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    set_text_frame(shape, text, font_size, font_color, bold, align)
    return shape

def add_rounded_rect(slide, left, top, width, height, color, text="", font_size=12, font_color=None):
    """添加圆角矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    set_shape_fill(shape, color)
    if font_color is None:
        font_color = COLORS['white']
    set_text_frame(shape, text, font_size, font_color, True, PP_ALIGN.CENTER)
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return shape

def add_rect(slide, left, top, width, height, color, text="", font_size=12, font_color=None):
    """添加矩形"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    set_shape_fill(shape, color)
    if font_color and text:
        set_text_frame(shape, text, font_size, font_color, True, PP_ALIGN.CENTER)
    return shape

def create_cover_slide(prs):
    """第1页：封面"""
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)

    # 深蓝背景
    background = add_rect(slide, 0, 0, 10, 7.5, COLORS['primary'])

    # 主标题
    add_text_box(slide, 0.5, 2.5, 9, 1, "游戏策划需求 Skill",
                 font_size=44, font_color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

    # 副标题
    add_text_box(slide, 0.5, 3.5, 9, 0.8, "从模糊想法到可交付 PRD + HTML 原型",
                 font_size=24, font_color=COLORS['secondary'], align=PP_ALIGN.CENTER)

    # 作者信息
    add_text_box(slide, 0.5, 5.5, 9, 0.5, "疏白 · 游戏策划",
                 font_size=18, font_color=COLORS['white'], align=PP_ALIGN.CENTER)

    # 日期
    add_text_box(slide, 0.5, 6, 9, 0.5, "2026.05.13",
                 font_size=14, font_color=COLORS['secondary'], align=PP_ALIGN.CENTER)

def create_pain_point_slide(prs):
    """第2页：痛点陈述"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # 浅灰背景
    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    # 标题
    add_text_box(slide, 0.5, 0.4, 9, 0.7, "🤔 策划需求的困境",
                 font_size=32, font_color=COLORS['primary'], bold=True, align=PP_ALIGN.LEFT)

    # 四个痛点卡片
    pains = [
        ("❓", "需求不清", "想法模糊，无从下手"),
        ("📝", "文档难写", "改了8遍还在返工"),
        ("🔄", "反复改稿", "方向不明，沟通成本高"),
        ("⏰", "时间紧迫", "deadline追着跑"),
    ]

    for i, (icon, title, desc) in enumerate(pains):
        left = 0.5 + i * 2.3
        # 卡片背景
        card = add_rounded_rect(slide, left, 1.4, 2.1, 2.2, COLORS['white'])
        card.line.color.rgb = COLORS['secondary']

        # 图标
        add_text_box(slide, left, 1.5, 2.1, 0.5, icon,
                     font_size=28, align=PP_ALIGN.CENTER)
        # 标题
        add_text_box(slide, left, 2.1, 2.1, 0.4, title,
                     font_size=16, font_color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)
        # 描述
        add_text_box(slide, left, 2.5, 2.1, 0.8, desc,
                     font_size=12, font_color=COLORS['text'], align=PP_ALIGN.CENTER)

    # 引用语
    quote_bg = add_rounded_rect(slide, 0.5, 4, 9, 2.5, COLORS['light_blue'])
    quote_bg.line.color.rgb = COLORS['secondary']

    quotes = [
        '"我想做一个公会系统，但不知道从哪开始"',
        '"写了三版需求文档，都被开发打回来重写"',
        '"设计稿改了8次，还在纠结交互"',
    ]

    for i, q in enumerate(quotes):
        add_text_box(slide, 1, 4.3 + i * 0.7, 8, 0.6, q,
                     font_size=16, font_color=COLORS['text'], align=PP_ALIGN.LEFT)

def create_workflow_overview(prs):
    """第3页：七步工作流全景图（详细版）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    # 标题
    add_text_box(slide, 0.5, 0.3, 9, 0.5, "💡 我们的方案：七步引导式产出",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 顶部：七步流程简图
    steps = [
        ("1", "需求\n采集"),
        ("2", "数据\n需求"),
        ("3", "方案\n设计"),
        ("4", "产出\n对齐"),
        ("5", "输出\n产物"),
        ("6", "设计\n集成"),
        ("7", "版本\n管理"),
    ]

    for i, (num, title) in enumerate(steps):
        left = 0.4 + i * 1.35
        # 圆形编号
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(left + 0.3), Inches(0.9), Inches(0.55), Inches(0.55)
        )
        set_shape_fill(circle, COLORS['primary'])
        set_text_frame(circle, num, 18, COLORS['white'], True, PP_ALIGN.CENTER)

        # 标题
        add_text_box(slide, left, 1.5, 1.15, 0.7, title,
                     font_size=11, font_color=COLORS['text'], bold=True, align=PP_ALIGN.CENTER)

        # 箭头（除了最后一个）
        if i < 6:
            add_text_box(slide, left + 1, 0.95, 0.4, 0.5, "→",
                         font_size=16, font_color=COLORS['secondary'], align=PP_ALIGN.CENTER)

    # 分隔线
    divider = add_rect(slide, 0.5, 2.25, 9, 0.03, COLORS['secondary'])

    # 中部：每一步详细说明
    step_details = [
        ("📝 Step 1 · 需求采集", '"7维度评估 + 3轮追问 → 理解确认"'),
        ("📊 Step 2 · 数据需求", '"现状数据 + 对比数据 + 验证数据 → 可行性判断"'),
        ("💡 Step 3 · 方案设计", '"瓶颈分析 + 方案清单 → 决策采纳"'),
        ("🎯 Step 4 · 产出对齐", '"PRD结构 + 原型需求 + 流程图需求 → 方向确认"'),
        ("📦 Step 5 · 输出产物", '"PRD文档 + HTML原型 + Mermaid流程图"'),
        ("🎨 Step 6 · Figma设计集成", '"设计稿截图 → 嵌入PRD文档"'),
        ("📋 Step 7 · 版本管理", '"CHANGELOG + 版本隔离 + 变更追溯"'),
    ]

    for i, (title, desc) in enumerate(step_details):
        top = 2.4 + i * 0.5
        row = i // 2
        col = i % 2
        left = 0.5 + col * 4.7

        # 编号圆点
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(left), Inches(top + 0.05), Inches(0.2), Inches(0.2)
        )
        set_shape_fill(dot, COLORS['secondary'])

        # 标题
        add_text_box(slide, left + 0.3, top, 2.5, 0.4, title,
                     font_size=12, font_color=COLORS['primary'], bold=True)
        # 描述
        add_text_box(slide, left + 2.8, top, 2, 0.4, desc,
                     font_size=10, font_color=COLORS['text'])

    # 底部：产出物区域
    output_bg = add_rounded_rect(slide, 0.5, 6.1, 9, 1.2, COLORS['light_blue'])
    output_bg.line.color.rgb = COLORS['primary']

    add_text_box(slide, 0.8, 6.2, 8, 0.4, "📦 最终产出物",
                 font_size=14, font_color=COLORS['primary'], bold=True)

    outputs = ["PRD 文档", "HTML 原型", "Mermaid 流程图", "设计稿"]
    for i, o in enumerate(outputs):
        left = 0.8 + i * 2.2
        item = add_rounded_rect(slide, left, 6.6, 2, 0.5, COLORS['secondary'])
        set_text_frame(item, o, 13, COLORS['white'], True, PP_ALIGN.CENTER)

def create_value_slide(prs):
    """第4页：核心价值"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.4, 9, 0.7, "⭐ 为什么用这套工作流？",
                 font_size=32, font_color=COLORS['primary'], bold=True)

    values = [
        ("🎯", "精准", "从需求到原型\n只需7步"),
        ("🚀", "高效", "减少沟通成本\n减少返工"),
        ("🔄", "迭代", "快速验证想法\n边做边改"),
    ]

    for i, (icon, title, desc) in enumerate(values):
        left = 0.5 + i * 3

        # 卡片
        card = add_rounded_rect(slide, left, 1.5, 2.8, 3, COLORS['white'])
        card.line.color.rgb = COLORS['secondary']

        # 图标
        add_text_box(slide, left, 1.7, 2.8, 0.8, icon,
                     font_size=36, align=PP_ALIGN.CENTER)

        # 标题
        add_text_box(slide, left, 2.5, 2.8, 0.6, title,
                     font_size=24, font_color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)

        # 描述
        add_text_box(slide, left, 3.2, 2.8, 1.2, desc,
                     font_size=14, font_color=COLORS['text'], align=PP_ALIGN.CENTER)

    # 底部总结
    add_text_box(slide, 0.5, 5.5, 9, 1, "告别「想到哪写到哪」，实现「按步骤高效产出」",
                 font_size=18, font_color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)

def create_step1_slide(prs):
    """第5页：Step 1 需求采集"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "📝 Step 1：需求采集 — 七维度评估",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 七维度
    dimensions = [
        ("#1", "用户与场景", "谁用？什么情况？"),
        ("#2", "核心目标", "解决什么问题？"),
        ("#3", "业务流程", "主流程是什么？"),
        ("#4", "边界条件", "异常情况有哪些？"),
        ("#5", "数据依赖", "需要哪些数据？"),
        ("#6", "系统影响", "影响哪些模块？"),
        ("#7", "优先级", "MVP 范围是什么？"),
    ]

    for i, (num, title, desc) in enumerate(dimensions):
        row = i // 4
        col = i % 4
        left = 0.5 + col * 2.3
        top = 1.1 + row * 1.2

        # 编号圆
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(left), Inches(top), Inches(0.5), Inches(0.5)
        )
        set_shape_fill(circle, COLORS['secondary'])
        set_text_frame(circle, num, 12, COLORS['white'], True, PP_ALIGN.CENTER)

        # 标题
        add_text_box(slide, left + 0.6, top, 1.6, 0.3, title,
                     font_size=14, font_color=COLORS['primary'], bold=True)
        # 描述
        add_text_box(slide, left + 0.6, top + 0.3, 1.6, 0.4, desc,
                     font_size=11, font_color=COLORS['text'])

    # 3轮追问原则
    principle_bg = add_rounded_rect(slide, 0.5, 4, 9, 2.8, COLORS['light_blue'])
    principle_bg.line.color.rgb = COLORS['primary']

    add_text_box(slide, 0.8, 4.2, 8, 0.5, "🔍 3轮追问原则",
                 font_size=18, font_color=COLORS['primary'], bold=True)

    principles = [
        "第1轮：澄清核心概念和业务目标",
        "第2轮：追问边界条件和异常流",
        "第3轮：确认优先级和依赖关系",
    ]

    for i, p in enumerate(principles):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(0.8), Inches(4.8 + i * 0.5), Inches(0.4), Inches(0.4)
        )
        set_shape_fill(circle, COLORS['accent'])
        set_text_frame(circle, str(i + 1), 14, COLORS['white'], True, PP_ALIGN.CENTER)
        add_text_box(slide, 1.3, 4.8 + i * 0.5, 7, 0.4, p,
                     font_size=14, font_color=COLORS['text'])

    add_text_box(slide, 0.8, 6.2, 8, 0.4, "✓ 最后复述理解，等待确认",
                 font_size=14, font_color=COLORS['secondary'], bold=True)

def create_step2_slide(prs):
    """第6页：Step 2 数据需求"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "📊 Step 2：数据需求 — 先问数据，再定方案",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 四象限图
    quadrants = [
        ("数据充足 ✓", "可推进方案", COLORS['secondary'], 0.5, 1.2),
        ("数据不足+可获取", "先补数据", COLORS['primary'], 3.5, 1.2),
        ("数据不足+不可获取", "方案搁置", COLORS['text'], 0.5, 3.2),
        ("无法验证", "方案放弃 ✗", COLORS['accent'], 3.5, 3.2),
    ]

    add_text_box(slide, 0.5, 1, 4, 0.3, "数据需求评估四象限",
                 font_size=14, font_color=COLORS['primary'], bold=True)

    for title, desc, color, left, top in quadrants:
        card = add_rounded_rect(slide, left, top, 2.8, 1.8, color)
        set_text_frame(card, f"{title}\n\n{desc}", 14, COLORS['white'], True, PP_ALIGN.CENTER)

    # 文档结构
    add_text_box(slide, 5.5, 1, 4, 0.4, "📄 数据需求文档结构",
                 font_size=16, font_color=COLORS['primary'], bold=True)

    structure = [
        "1. 方案所需数据",
        "2. 现有数据",
        "3. 数据缺口",
        "4. 获取计划",
        "5. 结论",
    ]

    for i, s in enumerate(structure):
        left = 5.5 if i < 3 else 7.5
        top = 1.5 + (i % 3) * 0.6
        item = add_rounded_rect(slide, left, top, 1.8, 0.45, COLORS['secondary'])
        set_text_frame(item, s, 12, COLORS['white'], True, PP_ALIGN.CENTER)

    # 核心原则
    principle_bg = add_rounded_rect(slide, 5.5, 3.5, 4, 1.8, COLORS['light_blue'])
    principle_bg.line.color.rgb = COLORS['accent']

    add_text_box(slide, 5.7, 3.7, 3.6, 0.4, "⚠️ 核心原则",
                 font_size=14, font_color=COLORS['accent'], bold=True)
    add_text_box(slide, 5.7, 4.2, 3.6, 1, "数据不充分\n不做方案\n\n方案不确定\n不写 PRD",
                 font_size=14, font_color=COLORS['text'])

def create_step3_slide(prs):
    """第7页：Step 3 方案设计"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "💡 Step 3：方案设计 — 方案不确定，不写 PRD",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 三模块
    modules = [
        ("1. 现状分析", ["瓶颈识别", "机会点"], 0.5),
        ("2. 方案清单", ["方案A/B/C", "预期收益", "实施难度", "依赖条件"], 3.5),
        ("3. 用户决策", ["采用", "不采纳", "修改后采纳"], 6.5),
    ]

    for title, items, left in modules:
        card = add_rounded_rect(slide, left, 1.2, 2.8, 3.5, COLORS['white'])
        card.line.color.rgb = COLORS['secondary']

        add_text_box(slide, left + 0.2, 1.4, 2.4, 0.5, title,
                     font_size=16, font_color=COLORS['primary'], bold=True)

        for i, item in enumerate(items):
            add_text_box(slide, left + 0.2, 2 + i * 0.5, 2.4, 0.4, f"• {item}",
                         font_size=13, font_color=COLORS['text'])

    # 核心原则
    principle_bg = add_rounded_rect(slide, 0.5, 5, 9, 1.8, COLORS['light_blue'])
    principle_bg.line.color.rgb = COLORS['accent']

    add_text_box(slide, 0.8, 5.2, 8, 0.5, "⚠️ 核心原则",
                 font_size=18, font_color=COLORS['accent'], bold=True)
    add_text_box(slide, 0.8, 5.8, 8, 1, "方案不清楚不往下走，数据不充分不做方案",
                 font_size=20, font_color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)

def create_step35_slide(prs):
    """第8页：Step 3.5 边界检验"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "🛡️ Step 3.5：边界检验 — 30项场景兜底",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 六大类型
    types = [
        ("资格边界", "5项", "用户升级/转官阶等"),
        ("时间边界", "5项", "中断登录/周期边界等"),
        ("道具边界", "5项", "背包已满/叠加等"),
        ("推送边界", "5项", "推送权限/token失效等"),
        ("技术边界", "5项", "接口超时/防重复等"),
        ("业务边界", "5项", "付费升级/账号注销等"),
    ]

    for i, (title, count, examples) in enumerate(types):
        row = i // 3
        col = i % 3
        left = 0.5 + col * 3.1
        top = 1.2 + row * 1.5

        card = add_rounded_rect(slide, left, top, 2.9, 1.3, COLORS['white'])
        card.line.color.rgb = COLORS['secondary']

        add_text_box(slide, left + 0.1, top + 0.1, 2.7, 0.4, title,
                     font_size=14, font_color=COLORS['primary'], bold=True)
        add_text_box(slide, left + 0.1, top + 0.5, 1, 0.4, count,
                     font_size=12, font_color=COLORS['accent'], bold=True)
        add_text_box(slide, left + 0.1, top + 0.9, 2.7, 0.3, examples,
                     font_size=10, font_color=COLORS['text'])

    # 优先级
    priority_bg = add_rounded_rect(slide, 0.5, 4.5, 9, 1.5, COLORS['light_blue'])
    priority_bg.line.color.rgb = COLORS['primary']

    add_text_box(slide, 0.8, 4.7, 8, 0.4, "优先级标注",
                 font_size=16, font_color=COLORS['primary'], bold=True)

    priorities = [
        ("🔴 P0 必须处理", "可能影响核心指标或引发客诉"),
        ("🟡 P1 建议处理", "影响用户体验但有规避方案"),
        ("⚪ P2 后期优化", "极端低频场景"),
    ]

    for i, (p, desc) in enumerate(priorities):
        left = 0.8 + i * 3
        add_text_box(slide, left, 5.3, 2.8, 0.5, p,
                     font_size=14, font_color=COLORS['text'], bold=True)

def create_step4_slide(prs):
    """第9页：Step 4 产出对齐"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "🎯 Step 4：产出对齐 — 先问再定，不做无用功",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 产出确认清单
    confirmations = [
        ("1️⃣", "PRD 结构确认", "需要几个 Tab？每个 Tab 放什么内容？"),
        ("2️⃣", "原型需求确认", "5轮追问确保原型方向正确"),
        ("3️⃣", "流程图需求", "复杂分支逻辑才画，简单流程略过"),
    ]

    for i, (icon, title, desc) in enumerate(confirmations):
        left = 0.5
        top = 1.2 + i * 1.8

        card = add_rounded_rect(slide, left, top, 9, 1.6, COLORS['white'])
        card.line.color.rgb = COLORS['secondary']

        add_text_box(slide, left + 0.3, top + 0.2, 1, 0.5, icon,
                     font_size=28, align=PP_ALIGN.CENTER)
        add_text_box(slide, left + 1.5, top + 0.2, 7, 0.5, title,
                     font_size=18, font_color=COLORS['primary'], bold=True)
        add_text_box(slide, left + 1.5, top + 0.7, 7, 0.8, desc,
                     font_size=14, font_color=COLORS['text'])

    # 原型5轮追问
    add_text_box(slide, 0.5, 5.8, 9, 0.4, "🔍 原型需求5轮追问：",
                 font_size=14, font_color=COLORS['accent'], bold=True)
    add_text_box(slide, 0.5, 6.2, 9, 1, "这个方案涉及几个界面？→ 每个界面的主要/次级内容？→ 用户的操作和状态变化？→ MVP范围？→ 同一内容需要几种布局？",
                 font_size=12, font_color=COLORS['text'])

def create_step5_slide(prs):
    """第10页：Step 5 输出产物"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "📦 Step 5：输出产物 — 标准化、可交付",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # 三类产出物
    outputs = [
        ("PRD 文档", ["prd_vX.X.html (引用外部)", "prd_vX.X_单文件版 (内联附件)"], COLORS['primary']),
        ("HTML 原型", ["prototype_vX.X.html", "可交互·可预览·可分享·可迭代"], COLORS['secondary']),
        ("Mermaid 流程图", ["flowchart_vX.X.md", "嵌入 PRD 文档"], COLORS['text']),
    ]

    for i, (title, items, color) in enumerate(outputs):
        left = 0.5 + i * 3.1
        top = 1.2

        card = add_rounded_rect(slide, left, top, 2.9, 3.5, color)

        add_text_box(slide, left + 0.2, top + 0.3, 2.5, 0.5, title,
                     font_size=16, font_color=COLORS['white'], bold=True)

        for j, item in enumerate(items):
            add_text_box(slide, left + 0.2, top + 1 + j * 0.8, 2.5, 0.7, item,
                         font_size=12, font_color=COLORS['white'])

    # 技术栈
    tech_bg = add_rounded_rect(slide, 0.5, 5, 9, 1.5, COLORS['light_blue'])
    tech_bg.line.color.rgb = COLORS['primary']

    add_text_box(slide, 0.8, 5.2, 8, 0.5, "⚡ 技术栈",
                 font_size=16, font_color=COLORS['primary'], bold=True)

    techs = ["Tailwind CSS", "Mermaid.js", "Chart.js"]
    for i, tech in enumerate(techs):
        left = 1 + i * 2.5
        item = add_rounded_rect(slide, left, 5.8, 2.2, 0.5, COLORS['secondary'])
        set_text_frame(item, tech, 14, COLORS['white'], True, PP_ALIGN.CENTER)

def create_step6_slide(prs):
    """第11页：Step 6 Figma 设计稿集成"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "🎨 Step 6：Figma 设计稿集成 — 视觉落地",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # PRD 主体框
    prd_bg = add_rounded_rect(slide, 1.5, 1.3, 7, 4.5, COLORS['white'])
    prd_bg.line.color.rgb = COLORS['primary']

    add_text_box(slide, 1.8, 1.5, 6.4, 0.5, "PRD 文档主体",
                 font_size=16, font_color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)

    # 设计稿截图区域
    for i in range(3):
        for j in range(2):
            img_placeholder = add_rounded_rect(
                slide,
                1.8 + j * 3.1,
                2.2 + i * 1.2,
                2.9, 1,
                COLORS['light_blue']
            )
            set_text_frame(img_placeholder, f"[设计稿截图 {i+1}-{j+1}]",
                           12, COLORS['text'], align=PP_ALIGN.CENTER)

    # AI 标注
    ai_tag = add_rounded_rect(slide, 7, 5.3, 1.8, 0.4, COLORS['accent'])
    set_text_frame(ai_tag, "↑ AI 自动截图", 10, COLORS['white'], True, PP_ALIGN.CENTER)

    # 设计规范
    add_text_box(slide, 0.5, 6.2, 9, 0.8,
                 "📐 设计规范：十周年画布尺寸 · 安全边距 · 弹窗按钮动效",
                 font_size=14, font_color=COLORS['secondary'], bold=True, align=PP_ALIGN.CENTER)

def create_step7_slide(prs):
    """第12页：Step 7 版本管理"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_rect(slide, 0, 0, 10, 7.5, COLORS['background'])

    add_text_box(slide, 0.5, 0.3, 9, 0.6, "📋 Step 7：版本管理 — 变更可追溯，版本不丢失",
                 font_size=28, font_color=COLORS['primary'], bold=True)

    # CHANGELOG 结构
    changelog_bg = add_rounded_rect(slide, 0.5, 1.2, 5.5, 4, COLORS['white'])
    changelog_bg.line.color.rgb = COLORS['secondary']

    add_text_box(slide, 0.8, 1.4, 5, 0.5, "CHANGELOG.md 结构",
                 font_size=16, font_color=COLORS['primary'], bold=True)

    changelog = """v1.0 (2026-05-13)
├─ 新增：初始版本
├─ 修改：奖励数值调整
└─ 修复：边界条件补充

v1.1 (2026-05-15)
├─ 新增：预流失标签体系
└─ 修改：召回链路流程优化"""

    add_text_box(slide, 0.8, 2, 5, 3, changelog,
                 font_size=13, font_color=COLORS['text'])

    # 版本隔离原则
    principle_bg = add_rounded_rect(slide, 6.2, 1.2, 3.3, 4, COLORS['light_blue'])
    principle_bg.line.color.rgb = COLORS['accent']

    add_text_box(slide, 6.4, 1.4, 3, 0.5, "⚠️ 版本隔离原则",
                 font_size=14, font_color=COLORS['accent'], bold=True)

    principles = [
        "每个版本物理隔离",
        "不覆盖历史版本",
        "原型每改一处",
        "PRD 同步更新一处",
    ]

    for i, p in enumerate(principles):
        add_text_box(slide, 6.4, 2 + i * 0.6, 3, 0.5, f"• {p}",
                     font_size=13, font_color=COLORS['text'])

def create_ending_slide(prs):
    """第13页：结尾 + 口诀"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # 深蓝背景
    add_rect(slide, 0, 0, 10, 7.5, COLORS['primary'])

    # CTA
    add_text_box(slide, 0.5, 0.8, 9, 0.8, "🚀 开始你的第一个需求",
                 font_size=36, font_color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

    # 口诀
    add_text_box(slide, 0.5, 1.8, 9, 0.5, "🎯 七步口诀",
                 font_size=20, font_color=COLORS['secondary'], bold=True, align=PP_ALIGN.CENTER)

    口诀_list = [
        "先问清，再查数据",
        "数据不够，先补缺口",
        "数据够了，再想方案",
        "方案不确定，不写 PRD",
        "产出方向，先问再定",
        "边界检验，漏洞补掉",
        "原型输出，多种布局",
        "设计稿放，集成 PRD",
        "版本管理，及时记录",
    ]

    # 两列布局
    for i, line in enumerate(口诀_list):
        col = i // 5
        row = i % 5
        left = 1 + col * 4.5
        top = 2.5 + row * 0.45

        num_circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(left), Inches(top), Inches(0.35), Inches(0.35)
        )
        set_shape_fill(num_circle, COLORS['secondary'])
        set_text_frame(num_circle, str(i + 1), 11, COLORS['white'], True, PP_ALIGN.CENTER)

        add_text_box(slide, left + 0.45, top, 3.8, 0.4, line,
                     font_size=14, font_color=COLORS['white'])

    # 底部信息
    add_text_box(slide, 0.5, 5.8, 9, 0.5, '"只需要描述你的想法，剩余交给我们"',
                 font_size=16, font_color=COLORS['secondary'], align=PP_ALIGN.CENTER)

    add_text_box(slide, 0.5, 6.5, 9, 0.5, "疏白 · 游戏策划",
                 font_size=18, font_color=COLORS['white'], align=PP_ALIGN.CENTER)

def main():
    # 创建演示文稿
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 按顺序创建所有幻灯片
    create_cover_slide(prs)           # 1. 封面
    create_pain_point_slide(prs)       # 2. 痛点
    create_workflow_overview(prs)      # 3. 全景图
    create_value_slide(prs)            # 4. 价值
    create_step1_slide(prs)            # 5. Step 1
    create_step2_slide(prs)            # 6. Step 2
    create_step3_slide(prs)            # 7. Step 3
    create_step35_slide(prs)           # 8. Step 3.5
    create_step4_slide(prs)            # 9. Step 4
    create_step5_slide(prs)            # 10. Step 5
    create_step6_slide(prs)            # 11. Step 6
    create_step7_slide(prs)            # 12. Step 7
    create_ending_slide(prs)           # 13. 结尾

    # 保存文件
    output_path = "c:/Users/zhangfan/my-openspec-project/游戏策划需求Skill_介绍PPT_v2.pptx"
    prs.save(output_path)
    print(f"PPT 已保存至: {output_path}")

if __name__ == "__main__":
    main()