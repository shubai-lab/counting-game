# -*- coding: utf-8 -*-
import openpyxl, os, sys
sys.stdout.reconfigure(encoding='utf-8')
path = r'C:\Users\zhangfan\Downloads\拉活功能分析-预流失干预\预流失干预数据分析报告_20260428_v2.xlsx'
wb = openpyxl.load_workbook(path)
print('所有Sheet:', wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'\n=== {sn} ===')
    for row in ws.iter_rows(values_only=True):
        print(row)
