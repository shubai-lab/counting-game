# -*- coding: utf-8 -*-
import pandas as pd
from openpyxl import load_workbook
import json

file_path = 'C:/Users/zhangfan/Desktop/武将品质划分_ 改版.xlsx'

# 使用 openpyxl 直接读取
wb = load_workbook(file_path, data_only=True)

result = {
    'sheets': wb.sheetnames,
    'content': {}
}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    data = []
    for row in ws.iter_rows(values_only=True):
        data.append([str(cell) if cell is not None else None for cell in row])

    result['content'][sheet_name] = {
        'rows': ws.max_row,
        'cols': ws.max_column,
        'data': data
    }

# 保存为 UTF-8 编码的 JSON
with open('excel_content_utf8.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=2)

print("Done. Saved to excel_content_utf8.json")
